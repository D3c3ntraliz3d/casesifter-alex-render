from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from app.engines import finding, timed_module
from app.ingestion import load_image, open_pdf_reader, render_pdf_pages
from app.models import Finding, FontUsage, ModuleResult

DATE_PATTERNS = [
    re.compile(r"\b(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b"),
    re.compile(
        r"\b((?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4})\b",
        re.I,
    ),
]


def analyze_ocr_text(path: Path, metadata_extras: dict[str, Any] | None = None) -> ModuleResult:
    return timed_module("ocr_text", lambda: _run(path, metadata_extras or {}))


def _run(path: Path, metadata_extras: dict[str, Any]) -> tuple[list[Finding], dict[str, Any]]:
    findings: list[Finding] = []
    fonts: list[FontUsage] = []
    extracted_pages: list[dict[str, Any]] = []
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text_pages, font_map, pdf_findings = _pdf_text(path)
        findings.extend(pdf_findings)
        extracted_pages = text_pages
        fonts = _fonts_from_map(font_map)
    elif suffix in {".jpg", ".jpeg", ".png", ".gif"}:
        text, ocr_findings = _ocr_image(load_image(path), page=1)
        findings.extend(ocr_findings)
        extracted_pages = [{"page": 1, "text": text, "source": "ocr"}]
    elif suffix == ".docx":
        extracted_pages, docx_findings = _docx_text(path)
        findings.extend(docx_findings)
    elif suffix == ".xlsx":
        extracted_pages = _xlsx_text(path)
    elif suffix == ".txt":
        extracted_pages = _plain_text(path)

    full_text = "\n".join(p.get("text") or "" for p in extracted_pages)
    extras: dict[str, Any] = {
        "pages": [
            {
                "page": p.get("page"),
                "source": p.get("source"),
                "char_count": len(p.get("text") or ""),
                "preview": (p.get("text") or "")[:800],
            }
            for p in extracted_pages
        ],
        "fonts": [f.model_dump() for f in fonts],
        "char_count": len(full_text),
        "full_text": full_text,
    }

    if suffix == ".pdf":
        findings.extend(_compare_ocr_vs_pdf_text(path, extracted_pages))

    doc_meta = (metadata_extras.get("metadata") or {}).get("document") or {}
    findings.extend(_date_mismatch(full_text, doc_meta))
    extras["font_models"] = fonts
    return findings, extras


def _pdf_text(path: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]], list[Finding]]:
    findings: list[Finding] = []
    reader = open_pdf_reader(path)
    pages: list[dict[str, Any]] = []
    font_map: dict[str, dict[str, Any]] = defaultdict(lambda: {"pages": set(), "sizes": set(), "sample": ""})
    invisible_hits = 0
    tiny_hits = 0

    for index, page in enumerate(reader.pages, start=1):
        runs: list[str] = []

        def visitor(text, cm, tm, font_dict, font_size):
            nonlocal tiny_hits
            snippet = (text or "").strip()
            if not snippet:
                return
            runs.append(snippet)
            name = "Unknown"
            if font_dict:
                name = str(font_dict.get("/BaseFont") or font_dict.get("/Name") or "Unknown")
            font_map[name]["pages"].add(index)
            try:
                size = float(font_size)
            except (TypeError, ValueError):
                size = 0.0
            if size:
                font_map[name]["sizes"].add(round(size, 1))
            if size and size < 0.6 and len(snippet) > 2:
                tiny_hits += 1
            if not font_map[name]["sample"]:
                font_map[name]["sample"] = snippet[:80]

        try:
            page.extract_text(visitor_text=visitor)
        except Exception:
            try:
                runs.append(page.extract_text() or "")
            except Exception:
                runs.append("")

        raw = ""
        try:
            raw = page.get_contents()
            if raw is not None:
                data = raw.get_data() if hasattr(raw, "get_data") else bytes(raw)
                if isinstance(data, bytes) and b"3 Tr" in data:
                    invisible_hits += 1
        except Exception:
            pass

        pages.append({"page": index, "text": " ".join(runs), "source": "pdf-text"})

    if invisible_hits:
        findings.append(
            finding(
                fid="txt-invisible-mode",
                category="text",
                severity="high",
                title="Invisible text rendering mode detected",
                explanation=(
                    f"{invisible_hits} page content stream(s) use PDF text render mode 3 (invisible). "
                    "This draws text that OCR/copy tools can still extract while remaining unseen on the page."
                ),
                why_suspicious="Hidden text is a common way to plant searchable content that does not match what a viewer sees.",
                confidence=84,
                evidence={"pages_with_invisible_mode": invisible_hits},
                recommended_action="Copy-paste the page text and compare it with the visible printout.",
            )
        )

    if tiny_hits:
        findings.append(
            finding(
                fid="txt-micro-text",
                category="text",
                severity="medium",
                title="Microscopic text runs found",
                explanation=(
                    f"{tiny_hits} text run(s) used a font size under 0.6 pt. That is far below readable print "
                    "and is a typical overlay trick."
                ),
                why_suspicious="Unreadable text that is still extractable does not belong on a normal court exhibit.",
                confidence=76,
                evidence={"tiny_runs": tiny_hits},
            )
        )

    # Font consistency: a letter/invoice using many unrelated fonts on one page.
    per_page_fonts: dict[int, set[str]] = defaultdict(set)
    for name, info in font_map.items():
        for page_no in info["pages"]:
            per_page_fonts[page_no].add(name)
    crowded = {p: fonts for p, fonts in per_page_fonts.items() if len(fonts) >= 10}
    if crowded:
        sample_page = next(iter(crowded))
        findings.append(
            finding(
                fid="txt-font-mix",
                category="text",
                severity="medium",
                title="Unusually high number of fonts on a single page",
                explanation=(
                    f"Page {sample_page} embeds {len(crowded[sample_page])} distinct fonts. "
                    "Native business documents typically use a handful. A collage of fonts can mean pages "
                    "or clauses were assembled from different sources."
                ),
                why_suspicious="Font explosion is a practical signal of cut-and-paste assembly.",
                confidence=64,
                page=sample_page,
                evidence={"fonts_on_page": sorted(crowded[sample_page])[:20]},
                recommended_action="Check whether letterhead, body, and amounts share a consistent typeface.",
            )
        )

    return pages, font_map, findings


def _fonts_from_map(font_map: dict[str, dict[str, Any]]) -> list[FontUsage]:
    fonts = []
    for name, info in sorted(font_map.items()):
        fonts.append(
            FontUsage(
                name=name,
                pages=sorted(info["pages"]),
                sizes=sorted(info["sizes"]),
                sample=info.get("sample") or None,
            )
        )
    return fonts


def _ocr_image(image, page: int | None) -> tuple[str, list[Finding]]:
    findings: list[Finding] = []
    try:
        import pytesseract

        text = pytesseract.image_to_string(image) or ""
        data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
        confs = [int(c) for c in data.get("conf", []) if str(c).isdigit() and int(c) >= 0]
        if confs and sum(1 for c in confs if c < 40) / max(len(confs), 1) > 0.45 and len(confs) > 20:
            findings.append(
                finding(
                    fid=f"txt-ocr-lowconf-{page or 0}",
                    category="text",
                    severity="info",
                    title="OCR confidence is uneven",
                    explanation="A large share of words decoded with low confidence. That can be a poor scan — or overlaid/erased text.",
                    why_suspicious="Erased or overwritten regions often confuse OCR more than a clean scan.",
                    confidence=40,
                    page=page,
                )
            )
        return text, findings
    except Exception as exc:
        return "", [
            finding(
                fid="txt-ocr-unavailable",
                category="text",
                severity="info",
                title="OCR engine not available on this server",
                explanation=(
                    "Tesseract is not installed or could not run "
                    f"({exc}). PDF embedded text was still analyzed when present."
                ),
                why_suspicious="Not an authenticity issue; it limits image-text checks.",
                confidence=0,
            )
        ]


def _compare_ocr_vs_pdf_text(path: Path, pdf_pages: list[dict[str, Any]]) -> list[Finding]:
    """If a page looks like a scan (little PDF text) we OCR a sample; if both exist, compare."""
    findings: list[Finding] = []
    text_chars = sum(len(p.get("text") or "") for p in pdf_pages)
    if text_chars > 80:
        return findings
    try:
        pages = render_pdf_pages(path, max_pages=min(3, len(pdf_pages) or 3))
    except Exception:
        return findings
    ocr_chunks = []
    for i, img in enumerate(pages, start=1):
        text, extra = _ocr_image(img, page=i)
        findings.extend([f for f in extra if f.id != "txt-ocr-unavailable"])
        ocr_chunks.append(text)
    ocr_text = " ".join(ocr_chunks).strip()
    if ocr_text and text_chars > 0:
        pdf_text = " ".join(p.get("text") or "" for p in pdf_pages)
        if _rough_mismatch(pdf_text, ocr_text):
            findings.append(
                finding(
                    fid="txt-ocr-mismatch",
                    category="text",
                    severity="high",
                    title="Visible (OCR) text does not match embedded PDF text",
                    explanation=(
                        "The file contains a small amount of extractable PDF text that does not align with "
                        "what optical character recognition reads from the page image. Hidden or overlay text is possible."
                    ),
                    why_suspicious="A mismatch between what you see and what you can copy is a classic overlay/redaction failure.",
                    confidence=78,
                    recommended_action="Copy the PDF text and compare it line-by-line with the rendered page.",
                )
            )
    return findings


def _rough_mismatch(a: str, b: str) -> bool:
    def tokens(s: str) -> set[str]:
        return {t.lower() for t in re.findall(r"[A-Za-z]{4,}", s)}

    ta, tb = tokens(a), tokens(b)
    if not ta or not tb:
        return False
    overlap = len(ta & tb) / max(1, min(len(ta), len(tb)))
    return overlap < 0.25 and min(len(ta), len(tb)) >= 8


def _docx_text(path: Path) -> tuple[list[dict[str, Any]], list[Finding]]:
    from docx import Document

    findings: list[Finding] = []
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    runs_fonts = []
    for p in doc.paragraphs:
        for run in p.runs:
            if run.text.strip():
                runs_fonts.append((run.font.name, run.text[:40]))
    names = {n for n, _ in runs_fonts if n}
    if len(names) >= 6:
        findings.append(
            finding(
                fid="txt-docx-fonts",
                category="text",
                severity="low",
                title="Many typefaces used in the Word document",
                explanation=f"Runs use fonts: {', '.join(sorted(names)[:12])}.",
                why_suspicious="Pasting from multiple sources often leaves mixed typefaces.",
                confidence=50,
                evidence={"fonts": sorted(names)},
            )
        )
    # White / hidden runs
    hidden = 0
    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.hidden or (run.font.color and getattr(run.font.color, "rgb", None) == "FFFFFF"):
                hidden += 1
    if hidden:
        findings.append(
            finding(
                fid="txt-docx-hidden",
                category="text",
                severity="high",
                title="Hidden or white-on-white Word text",
                explanation=f"{hidden} run(s) are marked hidden or white.",
                why_suspicious="Hidden runs can conceal clauses that still appear in some exports or metadata.",
                confidence=82,
            )
        )
    return [{"page": 1, "text": "\n".join(paragraphs), "source": "docx"}], findings


def _xlsx_text(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    pages = []
    for i, sheet in enumerate(wb.worksheets, start=1):
        rows = []
        for row in sheet.iter_rows(max_row=80, max_col=20, values_only=True):
            vals = [str(c) for c in row if c is not None]
            if vals:
                rows.append(" | ".join(vals))
        pages.append({"page": i, "text": "\n".join(rows), "source": f"xlsx:{sheet.title}"})
    wb.close()
    return pages


def _plain_text(path: Path) -> list[dict[str, Any]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    return [{"page": 1, "text": text, "source": "txt"}]


def _date_mismatch(text: str, doc_meta: dict[str, Any]) -> list[Finding]:
    created = doc_meta.get("created")
    if not created or not text:
        return []
    try:
        created_dt = datetime.fromisoformat(str(created).replace("Z", "+00:00"))
    except ValueError:
        return []
    mentioned = []
    for pattern in DATE_PATTERNS:
        mentioned.extend(pattern.findall(text[:8000]))
    if not mentioned:
        return []
    # If exhibit body is clearly years earlier than file creation, note it (scans of old docs are normal).
    years = []
    for item in mentioned[:15]:
        match = re.search(r"(19|20)\d{2}", str(item))
        if match:
            years.append(int(match.group(0)))
    if not years:
        return []
    body_year = min(years)
    if created_dt.year - body_year >= 3:
        return [
            finding(
                fid="txt-date-vs-file",
                category="text",
                severity="info",
                title="Document body date is years earlier than file creation",
                explanation=(
                    f"The text refers to year {body_year}, but the file’s embedded creation date is "
                    f"{created_dt.date().isoformat()}. This is expected for a later scan of an older paper original, "
                    "and unusual if the file is offered as a contemporaneous native electronic original."
                ),
                why_suspicious="A native file created years after the date it bears may be a recreation rather than the original.",
                confidence=45,
                evidence={"body_year": body_year, "file_created": created},
                recommended_action="Clarify whether this is a scan of a historical document or a native original.",
            )
        ]
    return []
